"""Read student assignments from an Excel workbook."""

import re
from io import BytesIO
from typing import TypedDict, cast

from openpyxl import load_workbook

HEADER_ALIASES = {
    "student_name": {"studentname", "nameofstudent", "nameofthestudent", "studentfullname", "fullname", "name"},
    "mentor_name": {"mentorname", "nameofmentor", "mentor", "teachername", "teacher", "guidename", "guide", "assignedmentor"},
    "prn": {"prn", "rollno", "rollnumber", "studentid", "enrollmentno", "enrollmentnumber"},
    "group_name": {"group", "groupname", "groupno", "groupnumber", "team", "teamname"},
    "project_name": {"project", "projectname", "projecttitle", "finalizedidea", "idea"},
    "year": {"year", "academicyear", "class"},
}


class MentorRosterRow(TypedDict):
    student_name: str
    mentor_name: str
    mentor_key: str
    prn: str
    group_name: str
    project_name: str
    year: str


def _text(value: object | None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _header(value: object | None) -> str:
    return re.sub(r"[^a-z0-9]", "", _text(value).lower())


def mentor_key(name: str) -> str:
    name = _text(name).casefold()
    title = r"^(?:professor|prof|dr|mrs|mr|ms)(?:\.\s*|\s+)"
    while re.match(title, name):
        name = re.sub(title, "", name)
    return " ".join(name.split())


def parse_mentor_roster(content: bytes, uploading_mentor: str) -> list[MentorRosterRow]:
    """Return rows from the first sheet with a Student Name header."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            columns: dict[str, int | None] | None = None
            for raw_row in rows:
                row = cast(tuple[object, ...], raw_row)
                columns = {
                    field: next((index for index, cell in enumerate(row) if _header(cell) in aliases), None)
                    for field, aliases in HEADER_ALIASES.items()
                }
                if columns["student_name"] is not None:
                    break
            if not columns or columns["student_name"] is None:
                continue

            def value(
                row: tuple[object, ...], field: str, *, columns: dict[str, int | None] = columns
            ) -> str:
                index = columns[field]
                return _text(row[index]) if index is not None and index < len(row) else ""

            students: list[MentorRosterRow] = []
            last_named_column = max(index for index in columns.values() if index is not None)
            current_group = ""
            current_mentor = uploading_mentor
            current_project = ""
            for raw_row in rows:
                row = cast(tuple[object, ...], raw_row)
                student_name = value(row, "student_name")
                group = value(row, "group_name")
                if group:
                    current_group = group
                    current_project = ""
                if columns["mentor_name"] is not None:
                    current_mentor = value(row, "mentor_name") or current_mentor
                current_project = value(row, "project_name") or current_project
                assigned_mentor = current_mentor
                if not assigned_mentor:
                    continue

                def add_student(
                    name: str,
                    prn: str,
                    *,
                    row: tuple[object, ...] = row,
                    students: list[MentorRosterRow] = students,
                    assigned_mentor: str = assigned_mentor,
                    current_group: str = current_group,
                    current_project: str = current_project,
                ) -> None:
                    students.append({
                        "student_name": name,
                        "mentor_name": assigned_mentor,
                        "mentor_key": mentor_key(assigned_mentor),
                        "prn": prn,
                        "group_name": current_group,
                        "project_name": current_project,
                        "year": value(row, "year"),
                    })
                    if len(students) > 5000:
                        raise ValueError("The sheet contains more than 5,000 students.")

                if student_name:
                    add_student(student_name, value(row, "prn"))

                # Some guide sheets place an additional PRN/name pair to the right.
                for index in range(last_named_column + 1, len(row) - 1):
                    extra_prn, extra_name = _text(row[index]), _text(row[index + 1])
                    if (re.fullmatch(r"[A-Za-z0-9-]{6,20}", extra_prn)
                            and any(char.isdigit() for char in extra_prn)
                            and extra_name and any(char.isalpha() for char in extra_name)):
                        add_student(extra_name, extra_prn)
            if not students:
                raise ValueError("The sheet has no student rows.")
            return students
        raise ValueError("Add a Student Name column to the Excel sheet.")
    finally:
        workbook.close()
